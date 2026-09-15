# -*- coding: utf-8 -*-
"""
טבלת חישוב חודשית ב-Excel עבור הנהלת החשבונות.

הדוח ב-Markdown נועד לקריאה. הגיליון הזה נועד לעבודה: רואת החשבון
ממיינת, מסננת, מתקנת סכום ורואה מיד איך הסיכום זז. לכן כל שורת סיכום
כאן היא נוסחה על גיליון הפירוט, ולא מספר שהודבק.

מבנה:
  סיכום   - טבלת החישוב. כל תא מחושב בנוסחה מגיליון הפירוט.
  פירוט   - שורה לכל מסמך, כולל שיעור ההכרה והחלק המוכר.
  לאימות  - רק השורות שהמנוע לא הכריע בהן. זו רשימת העבודה.
  לא נכלל - חשבונות עסקה ומטבע חוץ, עם הסיבה לאי-ההכללה.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- עיצוב

FONT_NAME = "Arial"
FMT_ILS = '#,##0.00;[Red]-#,##0.00'
FMT_PCT = '0%'

C_HEADER_BG = "1F3864"
C_HEADER_FG = "FFFFFF"
C_SECTION_BG = "D9E2F3"
C_BOTTOM_BG = "FFF2CC"
C_WARN_BG = "FCE4D6"
C_INPUT_FG = "0000FF"      # מספר שהוזן, לא מחושב
C_NOTE_FG = "808080"

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CATEGORY_LABELS = {
    "income": 'הכנסה - חשבונית מס',
    "expense": 'הוצאה מוכרת',
    "expense_no_vat": 'הוצאה ללא מע"מ',
    "expense_home": 'הוצאת בית - מוכרת חלקית',
    "credit": 'זיכוי',
    "proforma_out": 'חשבון עסקה שהופק',
    "proforma_in": 'חשבון עסקה שהתקבל',
}

EXCLUDED_REASONS = {
    "proforma_out": 'חשבון עסקה - אינו חשבונית מס, לא נכנס לדוח',
    "proforma_in": 'חשבון עסקה - אינו מזכה בניכוי תשומות',
}

DETAIL = "פירוט"

# עמודות גיליון הפירוט
COLS = [
    ("מסמך", 46),
    ("קוד", 16),
    ("סוג", 26),
    ("מטבע", 8),
    ("נטו", 14),
    ('מע"מ', 14),
    ('סה"כ', 14),
    ("שיעור הכרה", 11),
    ("נטו מוכר", 14),
    ('מע"מ מוכר', 14),
    ("אומדן", 8),
    ("הערה", 52),
]
# A מסמך  B קוד  C סוג  D מטבע  E נטו  F מע"מ  G סה"כ
# H שיעור  I נטו מוכר  J מע"מ מוכר  K אומדן  L הערה


def _rtl(ws):
    ws.sheet_view.rightToLeft = True


def _style_header(ws, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color=C_HEADER_FG)
        cell.fill = PatternFill("solid", fgColor=C_HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------- פירוט


def _sheet_detail(wb: Workbook, rows: list) -> int:
    """כותב את גיליון הפירוט ומחזיר את מספר השורה האחרונה שיש בה נתונים."""
    ws = wb.create_sheet(DETAIL)
    _rtl(ws)
    _widths(ws, [w for _, w in COLS])

    for i, (title, _) in enumerate(COLS, start=1):
        ws.cell(row=1, column=i, value=title)
    _style_header(ws, 1, len(COLS))

    r = 2
    for row in rows:
        rate = getattr(row, "recognized_rate", 1.0) or 1.0
        ws.cell(row=r, column=1, value=row.file)
        ws.cell(row=r, column=2, value=row.category)
        ws.cell(row=r, column=3, value=CATEGORY_LABELS.get(row.category, row.category))
        ws.cell(row=r, column=4, value=row.currency)
        ws.cell(row=r, column=5, value=round(row.net or 0.0, 2))
        ws.cell(row=r, column=6, value=round(row.vat or 0.0, 2))
        ws.cell(row=r, column=7, value=round(row.total or 0.0, 2))
        ws.cell(row=r, column=8, value=rate)
        # מראה את חישוב המנוע: בהוצאת בית ללא נטו מזוהה, הבסיס הוא הסכום המלא
        ws.cell(row=r, column=9,
                value=f'=IF($B{r}="expense_home",IF($E{r}>0,$E{r},$G{r})*$H{r},$E{r}*$H{r})')
        ws.cell(row=r, column=10, value=f'=$F{r}*$H{r}')
        ws.cell(row=r, column=11, value="כן" if row.estimated else "לא")
        ws.cell(row=r, column=12, value=row.note or "")
        r += 1

    last = r - 1
    for rr in range(2, r):
        for c in range(1, len(COLS) + 1):
            cell = ws.cell(row=rr, column=c)
            cell.font = Font(name=FONT_NAME, size=10,
                             color=C_INPUT_FG if c in (5, 6, 7, 8) else "000000")
            cell.border = BOX
            if c in (5, 6, 7, 9, 10):
                cell.number_format = FMT_ILS
            elif c == 8:
                cell.number_format = FMT_PCT
                cell.alignment = Alignment(horizontal="center")
            elif c in (4, 11):
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=(c == 12))
        if ws.cell(row=rr, column=11).value == "כן":
            ws.cell(row=rr, column=11).fill = PatternFill("solid", fgColor=C_WARN_BG)

    if last >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{last}"

    note = ws.cell(row=last + 2, column=1,
                   value='כחול = סכום שחולץ מהמסמך. שחור = מחושב בנוסחה. '
                         'תיקון בעמודות נטו/מע"מ/סה"כ מתעדכן מיד בגיליון הסיכום.')
    note.font = Font(name=FONT_NAME, size=9, italic=True, color=C_NOTE_FG)
    return last


# ---------------------------------------------------------------- סיכום


def _sumifs(col: str, code: str, last: int, currency: str = "ILS") -> str:
    rng = f"'{DETAIL}'!${col}$2:${col}${last}"
    return (f"SUMIFS({rng},'{DETAIL}'!$B$2:$B${last},\"{code}\","
            f"'{DETAIL}'!$D$2:$D${last},\"{currency}\")")


def _sheet_summary(wb: Workbook, totals, month_title: str, last: int,
                   home_rate: float, home_from: str, doc_count: int):
    ws = wb.create_sheet("סיכום", 0)
    _rtl(ws)
    _widths(ws, [42, 18, 18, 62])

    ws.merge_cells("A1:D1")
    t = ws.cell(row=1, column=1, value=f"טבלת חישוב חודשית - {month_title}")
    t.font = Font(name=FONT_NAME, size=14, bold=True, color=C_HEADER_BG)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:D2")
    s = ws.cell(row=2, column=1,
                value=f"עו\"ד ד\"ר אברהם ללום | {doc_count} מסמכים | "
                      f"כל הסכומים בשקלים, מחושבים מגיליון הפירוט")
    s.font = Font(name=FONT_NAME, size=9, italic=True, color=C_NOTE_FG)
    s.alignment = Alignment(horizontal="center")

    for i, h in enumerate(["סעיף", "נטו", 'מע"מ', "הערה"], start=1):
        ws.cell(row=4, column=i, value=h)
    _style_header(ws, 4, 4)

    def section(r, label):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = ws.cell(row=r, column=1, value=label)
        c.font = Font(name=FONT_NAME, size=11, bold=True, color=C_HEADER_BG)
        c.fill = PatternFill("solid", fgColor=C_SECTION_BG)
        c.alignment = Alignment(horizontal="right", vertical="center")
        for k in range(1, 5):
            ws.cell(row=r, column=k).border = BOX

    def line(r, label, net_f, vat_f, note="", bold=False, fill=None):
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=net_f)
        ws.cell(row=r, column=3, value=vat_f)
        ws.cell(row=r, column=4, value=note)
        for k in range(1, 5):
            cell = ws.cell(row=r, column=k)
            cell.border = BOX
            cell.font = Font(name=FONT_NAME, size=10 if not bold else 11, bold=bold)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            if k in (2, 3):
                cell.number_format = FMT_ILS
                cell.alignment = Alignment(horizontal="center")
            elif k == 4:
                cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")

    r = 5
    section(r, "הכנסות"); r += 1
    R_INC = r
    line(r, "חשבוניות מס שהופקו", f"={_sumifs('I','income',last)}",
         f"={_sumifs('J','income',last)}",
         'רק חשבוניות מס. חשבונות עסקה מופיעים בגיליון "לא נכלל".')
    r += 2

    section(r, "הוצאות"); r += 1
    R_EXP = r
    line(r, 'הוצאות מוכרות עם מע"מ',
         f"={_sumifs('I','expense',last)}-{_sumifs('I','credit',last)}",
         f"={_sumifs('J','expense',last)}-{_sumifs('J','credit',last)}",
         "בניכוי זיכויים")
    r += 1
    R_NOVAT = r
    line(r, 'הוצאות ללא מע"מ', f"={_sumifs('G','expense_no_vat',last)}", None,
         'ארנונה, אגרות, ביטוח, וקבלות שאינן חשבונית מס. מוכרות כהוצאה, לא מזכות בתשומות.')
    r += 2

    section(r, f"הוצאות בית - הכרה חלקית {int(home_rate*100)}%"); r += 1
    R_HOME_FULL = r
    line(r, "סכום מלא ששולם", f"={_sumifs('G','expense_home',last)}", None,
         "לשקיפות בלבד, לא נכנס לחישוב")
    r += 1
    R_HOME = r
    line(r, f"החלק המוכר ({int(home_rate*100)}%)",
         f"={_sumifs('I','expense_home',last)}",
         f"={_sumifs('J','expense_home',last)}",
         f"חשמל, מים וארנונה המשויכים לבית. חל מ-{home_from}.")
    r += 2

    section(r, "שורה תחתונה"); r += 1
    R_PROFIT = r
    line(r, 'רווח גולמי (לפני מע"מ)',
         f"=B{R_INC}-(B{R_EXP}+B{R_NOVAT}+B{R_HOME})", None,
         "הכנסות נטו פחות כל ההוצאות נטו, כולל אלה ללא תשומות",
         bold=True, fill=C_BOTTOM_BG)
    r += 1
    R_VAT = r
    line(r, 'מע"מ לתשלום', None,
         f"=C{R_INC}-(C{R_EXP}+C{R_HOME})",
         'מע"מ עסקאות פחות מע"מ תשומות. סכום שלילי = החזר.',
         bold=True, fill=C_BOTTOM_BG)
    r += 2

    # ------------------------------------------------ הנחות, גלויות לקורא
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    h = ws.cell(row=r, column=1, value="הנחות החישוב")
    h.font = Font(name=FONT_NAME, size=11, bold=True, color=C_HEADER_BG)
    h.fill = PatternFill("solid", fgColor=C_SECTION_BG)
    h.alignment = Alignment(horizontal="right")
    for k in range(1, 5):
        ws.cell(row=r, column=k).border = BOX
    r += 1

    assumptions = [
        (f"הכרה בהוצאות בית: {int(home_rate*100)}%",
         f"לפי הנחיית הלקוח, חל על חשמל, מים וארנונה מ-{home_from} ואילך. "
         "אותו שיעור מוחל על ההוצאה ועל מס התשומות."),
        ("בסיס הדיווח: מזומן",
         "המנוע סופר רק חשבונית מס קבלה. חשבונות עסקה שטרם שולמו אינם נספרים. "
         "אם הדיווח בפועל על בסיס מצטבר - יש לעדכן, ההפרש מהותי."),
        ("סכומים מסומנים כאומדן",
         'כל שורה שסומנה "כן" בעמודת האומדן בגיליון הפירוט חולצה בזיהוי טקסט '
         'ולא אומתה. ראו גיליון "לאימות".'),
        ("מקור הנתונים",
         "קבצי ה-PDF בתיקיית החודש בלבד. מסמך שלא הגיע לתיקייה אינו מופיע כאן, "
         "והמנוע אינו יכול לדעת שהוא חסר."),
    ]
    for label, text in assumptions:
        ws.cell(row=r, column=1, value=label)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.cell(row=r, column=2, value=text)
        ws.cell(row=r, column=1).font = Font(name=FONT_NAME, size=9, bold=True)
        ws.cell(row=r, column=2).font = Font(name=FONT_NAME, size=9)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="right", vertical="top",
                                                       wrap_text=True)
        for k in range(1, 5):
            ws.cell(row=r, column=k).border = BOX
        ws.row_dimensions[r].height = 30
        r += 1

    # מטבע חוץ, אם יש
    foreign = getattr(totals, "foreign", {}) or {}
    if foreign:
        r += 1
        ws.cell(row=r, column=1, value="הוצאות במטבע חוץ")
        ws.cell(row=r, column=1).font = Font(name=FONT_NAME, size=10, bold=True)
        r += 1
        for cur, amount in sorted(foreign.items()):
            ws.cell(row=r, column=1, value=cur)
            ws.cell(row=r, column=2, value=round(amount, 2))
            ws.cell(row=r, column=2).number_format = '#,##0.00'
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
            ws.cell(row=r, column=3,
                    value="לא הומר לשקל ולא נכנס לסיכום. נדרש שער יציג ליום החיוב.")
            for k in range(1, 5):
                ws.cell(row=r, column=k).font = Font(name=FONT_NAME, size=9)
                ws.cell(row=r, column=k).border = BOX
            r += 1

    return ws


# ---------------------------------------------------------------- לאימות


def _sheet_review(wb: Workbook, rows: list):
    flagged = [x for x in rows
               if getattr(x, "estimated", False)
               or (x.note and x.category not in ("proforma_out", "proforma_in"))]
    ws = wb.create_sheet("לאימות")
    _rtl(ws)
    _widths(ws, [46, 26, 14, 14, 14, 64])

    ws.merge_cells("A1:F1")
    t = ws.cell(row=1, column=1,
                value="שורות שהמנוע לא הכריע בהן. זו רשימת העבודה, לא רשימת שגיאות.")
    t.font = Font(name=FONT_NAME, size=10, bold=True, color=C_HEADER_BG)
    t.alignment = Alignment(horizontal="right")

    for i, h in enumerate(["מסמך", "סוג", "נטו", 'מע"מ', 'סה"כ', "מה דורש בדיקה"], start=1):
        ws.cell(row=3, column=i, value=h)
    _style_header(ws, 3, 6)

    r = 4
    for x in flagged:
        ws.cell(row=r, column=1, value=x.file)
        ws.cell(row=r, column=2, value=CATEGORY_LABELS.get(x.category, x.category))
        ws.cell(row=r, column=3, value=round(x.net or 0.0, 2))
        ws.cell(row=r, column=4, value=round(x.vat or 0.0, 2))
        ws.cell(row=r, column=5, value=round(x.total or 0.0, 2))
        reason = x.note or ""
        if getattr(x, "estimated", False):
            reason = ("הסכום חולץ בזיהוי טקסט ולא אומת. " + reason).strip()
        ws.cell(row=r, column=6, value=reason)
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.border = BOX
            if c in (3, 4, 5):
                cell.number_format = FMT_ILS
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        r += 1

    if not flagged:
        c = ws.cell(row=4, column=1, value="אין שורות לאימות. כל הסכומים חולצו מהמסמכים והוכרעו.")
        c.font = Font(name=FONT_NAME, size=10, italic=True, color="006100")
    return len(flagged)


# ---------------------------------------------------------------- לא נכלל


def _sheet_excluded(wb: Workbook, rows: list):
    excluded = [x for x in rows
                if x.category in ("proforma_out", "proforma_in") or x.currency != "ILS"]
    ws = wb.create_sheet("לא נכלל")
    _rtl(ws)
    _widths(ws, [46, 26, 10, 16, 64])

    ws.merge_cells("A1:E1")
    t = ws.cell(row=1, column=1,
                value="מסמכים שלא נכנסו לחישוב, והסיבה. חשבון עסקה שכבר שולם - נא לסמן.")
    t.font = Font(name=FONT_NAME, size=10, bold=True, color=C_HEADER_BG)
    t.alignment = Alignment(horizontal="right")

    for i, h in enumerate(["מסמך", "סוג", "מטבע", 'סה"כ', "סיבת אי-ההכללה"], start=1):
        ws.cell(row=3, column=i, value=h)
    _style_header(ws, 3, 5)

    r = 4
    for x in excluded:
        ws.cell(row=r, column=1, value=x.file)
        ws.cell(row=r, column=2, value=CATEGORY_LABELS.get(x.category, x.category))
        ws.cell(row=r, column=3, value=x.currency)
        ws.cell(row=r, column=4, value=round(x.total or 0.0, 2))
        reason = EXCLUDED_REASONS.get(x.category, "")
        if x.currency != "ILS":
            reason = (reason + " מטבע חוץ, נדרש שער יציג ליום החיוב.").strip()
        ws.cell(row=r, column=5, value=reason)
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.border = BOX
            if c == 4:
                cell.number_format = FMT_ILS
                cell.alignment = Alignment(horizontal="center")
            elif c == 3:
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        r += 1

    if excluded:
        ws.cell(row=r + 1, column=1, value='סה"כ')
        ws.cell(row=r + 1, column=4, value=f"=SUM(D4:D{r-1})")
        for c in (1, 4):
            cell = ws.cell(row=r + 1, column=c)
            cell.font = Font(name=FONT_NAME, size=10, bold=True)
            cell.fill = PatternFill("solid", fgColor=C_BOTTOM_BG)
            cell.border = BOX
        ws.cell(row=r + 1, column=4).number_format = FMT_ILS
    else:
        c = ws.cell(row=4, column=1, value="אין מסמכים שנותרו מחוץ לחישוב.")
        c.font = Font(name=FONT_NAME, size=10, italic=True, color="006100")
    return len(excluded)


# ---------------------------------------------------------------- API


def build_workbook(rows: list, totals, month_title: str, out_path: Path,
                   home_rate: float = 0.25, home_from: str = "2026-09") -> Path:
    """בונה את חוברת העבודה החודשית ושומר אותה ב-out_path."""
    wb = Workbook()
    wb.remove(wb.active)

    last = _sheet_detail(wb, rows)
    _sheet_summary(wb, totals, month_title, last, home_rate, home_from, len(rows))
    _sheet_review(wb, rows)
    _sheet_excluded(wb, rows)

    wb.active = 0
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def build_for_month(month: str, folder: Path, rows: list, totals,
                    home_rate: float = 0.25, home_from: str = "2026-09") -> Path:
    """נקודת הכניסה מהפייפליין. מחזיר את נתיב הקובץ שנשמר."""
    from ..accounting import _month_title
    out = Path(folder) / f"טבלת חישוב {month}.xlsx"
    return build_workbook(rows, totals, _month_title(month), out,
                          home_rate=home_rate, home_from=home_from)
